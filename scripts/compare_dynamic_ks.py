#!/usr/bin/env python
import argparse, os, pickle, random
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

def load_debug_data(debug_dir):
    dp = Path(debug_dir)
    if not dp.exists():
        raise FileNotFoundError(f"Directory not found: {debug_dir}")
    all_rec = []
    pkls = sorted(dp.glob("*.pkl"))
    if not pkls:
        raise FileNotFoundError(f"No .pkl files in {debug_dir}")
    print(f"Found {len(pkls)} pkl files")
    for pf in pkls:
        with open(pf, "rb") as f:
            recs = pickle.load(f)
            all_rec.extend(recs)
            print(f"  {pf.name}: {len(recs)} records")
    print(f"Total: {len(all_rec)}")
    return all_rec

def bucket_and_sample(records, max_per=30, seed=42):
    random.seed(seed)
    small  = [r for r in records if r["area"] < 655]
    medium = [r for r in records if 655 <= r["area"] < 9216]
    large  = [r for r in records if r["area"] >= 9216]
    print(f"Small (<655):    {len(small)}")
    print(f"Medium (655-9216):{len(medium)}")
    print(f"Large (>9216):    {len(large)}")
    sampled = []
    for name, bucket in [("small",small),("medium",medium),("large",large)]:
        if bucket:
            n = min(max_per, len(bucket))
            sampled.extend(random.sample(bucket, n))
            print(f"  Sampled {n} from {name}")
    sampled.sort(key=lambda r: r["area"])
    print(f"Total sampled: {len(sampled)}")
    return sampled

def generate_excel(records, output):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "dynamic_ks"
    H = Font(name="Arial", size=11, bold=True)
    HF = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    SH = Font(name="Arial", size=10, bold=True)
    DF = Font(name="Arial", size=10)
    TB = Border(left=Side(style="thin"),right=Side(style="thin"),
                top=Side(style="thin"),bottom=Side(style="thin"))
    CA = Alignment(horizontal="center", vertical="center")

    ws["B1"] = "真值框大小"
    ws["B1"].font = H; ws["B1"].fill = HF; ws["B1"].alignment = CA; ws["B1"].border = TB
    ws.merge_cells("C1:E1")
    ws["C1"] = "匹配正样本个数"
    ws["C1"].font = H; ws["C1"].fill = HF; ws["C1"].alignment = CA; ws["C1"].border = TB
    ws["B2"].border = TB
    ws["C2"] = "基于iou匹配"
    ws["C2"].font = SH
    ws["C2"].fill = PatternFill(start_color="E2EFDA",end_color="E2EFDA",fill_type="solid")
    ws["C2"].alignment = CA; ws["C2"].border = TB
    ws["D2"] = "本文方法"
    ws["D2"].font = SH
    ws["D2"].fill = PatternFill(start_color="E2EFDA",end_color="E2EFDA",fill_type="solid")
    ws["D2"].alignment = CA; ws["D2"].border = TB
    ws["E2"] = "\u57fa\u7840\u5206\u914d\u5668"
    ws["E2"].font = SH
    ws["E2"].fill = PatternFill(start_color="FCE4D6",end_color="FCE4D6",fill_type="solid")
    ws["E2"].alignment = CA; ws["E2"].border = TB

    for i, r in enumerate(records, start=3):
        ws.cell(row=i,column=2,value=r["area"]).font=DF
        ws.cell(row=i,column=2).alignment=CA; ws.cell(row=i,column=2).border=TB
        ws.cell(row=i,column=3,value=r["dynamic_ks_iou"]).font=DF
        ws.cell(row=i,column=3).alignment=CA; ws.cell(row=i,column=3).border=TB
        ws.cell(row=i,column=4,value=r["dynamic_ks"]).font=DF
        ws.cell(row=i,column=4).alignment=CA; ws.cell(row=i,column=4).border=TB
        ws.cell(row=i,column=5,value=r.get("basic_pos_count",0)).font=DF
        ws.cell(row=i,column=5).alignment=CA; ws.cell(row=i,column=5).border=TB

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16

    ro = len(records) + 5
    ws.cell(row=ro,column=2,value="统计").font=Font(name="Arial",size=11,bold=True)
    ws.merge_cells(f"C{ro}:E{ro}")
    ws.cell(row=ro,column=3,value="均值(iou/本文)").font=Font(name="Arial",size=10,bold=True)

    n = len(records)
    avg_i = sum(r["dynamic_ks_iou"] for r in records) / n
    avg_s = sum(r["dynamic_ks"] for r in records) / n
    ws.cell(row=ro+1,column=2,value="总体均值")
    ws.cell(row=ro+1,column=3,value=round(avg_i,2))
    ws.cell(row=ro+1,column=4,value=round(avg_s,2))
    avg_b = sum(r.get("basic_pos_count",0) for r in records) / n
    ws.cell(row=ro+1,column=5,value=round(avg_b,2))

    for j,(lbl,lo,hi) in enumerate([
        ("Small <655", 0, 655),
        ("Medium 655-9216", 655, 9216),
        ("Large >9216", 9216, float("inf")),
    ]):
        bucket = [r for r in records if lo <= r["area"] < hi]
        if bucket:
            b_n = len(bucket)
            ws.cell(row=ro+2+j,column=2,value=lbl)
            ws.cell(row=ro+2+j,column=3,value=round(sum(r["dynamic_ks_iou"] for r in bucket)/b_n,2))
            ws.cell(row=ro+2+j,column=4,value=round(sum(r["dynamic_ks"] for r in bucket)/b_n,2))
            ws.cell(row=ro+2+j,column=5,value=round(sum(r.get("basic_pos_count",0) for r in bucket)/b_n,2))

    wb.save(output)
    print(f"\nExcel saved: {output}")
    print(f"Rows: {n}")
    print(f"Avg k_iou: {avg_i:.2f}  Avg k_sa: {avg_s:.2f}  Avg basic: {avg_b:.2f}  Delta(sa-iou): {avg_s-avg_i:+.2f}")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--debug_dir", type=str, required=True)
    ap.add_argument("--output", type=str, default="compare_dynamic_ks.xlsx")
    ap.add_argument("--max_per_bucket", type=int, default=300)
    ap.add_argument("--no_sample", action="store_true")
    args = ap.parse_args()
    records = load_debug_data(args.debug_dir)
    if not args.no_sample:
        records = bucket_and_sample(records, max_per=args.max_per_bucket)
    generate_excel(records, args.output)

if __name__ == "__main__":
    main()
